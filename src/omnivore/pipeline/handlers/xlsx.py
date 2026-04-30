from __future__ import annotations

import io
from typing import ClassVar

import structlog

from omnivore.pipeline.context import BlobRef, IngestContext
from omnivore.pipeline.models import ExtractionResult, Fragment, StructuredTable

logger = structlog.get_logger(__name__)


class XlsxHandler:
    name: ClassVar[str] = "xlsx"
    version: ClassVar[str] = "1.0.0"
    accepts: ClassVar[tuple[str, ...]] = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
    )
    cost_class: ClassVar[str] = "io"
    timeout_seconds: ClassVar[int] = 120

    async def extract(self, blob: BlobRef, ctx: IngestContext) -> ExtractionResult:
        import openpyxl

        data = await ctx.read_blob()
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)

        result = ExtractionResult(
            document_id=ctx.document_id,
            source_handler=self.name,
            handler_version=self.version,
            metadata={"format": "xlsx", "sheet_count": len(wb.sheetnames), "sheets": wb.sheetnames},
        )

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_iter = list(ws.iter_rows(values_only=True))
            if not rows_iter:
                continue

            headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows_iter[0])]
            rows: list[dict] = []
            for raw_row in rows_iter[1:]:
                row = {headers[i]: (str(v) if v is not None else "") for i, v in enumerate(raw_row)}
                if any(row.values()):
                    rows.append(row)

            table = StructuredTable(
                name=sheet_name,
                headers=headers,
                rows=rows,
                sheet=sheet_name,
            )
            result.tables.append(table)

            row_rag_enabled = ctx.config.get("row_level_rag", False)
            if row_rag_enabled:
                for i, row in enumerate(rows[:500]):
                    content = " | ".join(f"{k}: {v}" for k, v in row.items() if v)
                    result.fragments.append(
                        Fragment(
                            kind="table_row",
                            content=content,
                            position={"row_index": i, "sheet": sheet_name},
                            table_lineage={"table_id": str(table.table_id), "row_idx": i},
                        )
                    )

        wb.close()
        logger.info(
            "xlsx.extracted",
            document_id=str(ctx.document_id),
            sheets=len(wb.sheetnames),
            tables=len(result.tables),
        )
        return result
