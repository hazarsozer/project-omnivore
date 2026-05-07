from __future__ import annotations

import io
from typing import ClassVar

import structlog

from omnivore.pipeline.context import BlobRef, IngestContext
from omnivore.pipeline.models import Block, ExtractionResult, Fragment, PagePosition, StructuredTable

logger = structlog.get_logger(__name__)

_HEADING_STYLES = {f"Heading {i}": i for i in range(1, 10)}


class DocxHandler:
    name: ClassVar[str] = "docx"
    version: ClassVar[str] = "1.0.0"
    accepts: ClassVar[tuple[str, ...]] = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    )
    cost_class: ClassVar[str] = "cpu"
    timeout_seconds: ClassVar[int] = 120

    async def extract(self, blob: BlobRef, ctx: IngestContext) -> ExtractionResult:
        from docx import Document

        data = await ctx.read_blob()
        doc = Document(io.BytesIO(data))
        result = ExtractionResult(
            document_id=ctx.document_id,
            source_handler=self.name,
            handler_version=self.version,
            metadata={"format": "docx"},
        )

        reading_order = 0
        heading_stack: list[tuple[int, str]] = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ""
            heading_level = _HEADING_STYLES.get(style_name)

            if heading_level is not None:
                heading_stack = [(lvl, t) for lvl, t in heading_stack if lvl < heading_level]
                heading_stack.append((heading_level, text))
                blk = Block(
                    kind="heading",
                    level=heading_level,
                    reading_order=reading_order,
                    text=text,
                )
            else:
                blk = Block(
                    kind="paragraph",
                    reading_order=reading_order,
                    text=text,
                )
                result.fragments.append(
                    Fragment(
                        kind="text",
                        content=text,
                        position=PagePosition(page=1),
                        source_block_ids=[blk.block_id],
                        heading_path=[t for _, t in heading_stack],
                    )
                )
            result.blocks.append(blk)
            reading_order += 1

        for t_idx, table in enumerate(doc.tables):
            if not table.rows:
                continue
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            rows = [
                {headers[i]: cell.text.strip() for i, cell in enumerate(row.cells)}
                for row in table.rows[1:]
            ]
            result.tables.append(
                StructuredTable(
                    name=f"table_{t_idx + 1}",
                    headers=headers,
                    rows=rows,
                )
            )

        logger.info(
            "docx.extracted",
            document_id=str(ctx.document_id),
            fragments=len(result.fragments),
            tables=len(result.tables),
        )
        return result
