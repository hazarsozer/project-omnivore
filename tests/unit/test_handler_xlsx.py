"""Unit tests for XlsxHandler — openpyxl builds in-memory workbooks, no disk I/O."""
from __future__ import annotations

import io
import uuid
from unittest.mock import AsyncMock, MagicMock

from omnivore.pipeline.context import BlobRef, IngestContext
from omnivore.pipeline.handlers.xlsx import XlsxHandler

# ---------------------------------------------------------------------------
# Fixture helpers
# ---------------------------------------------------------------------------

def _build_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Build an in-memory XLSX.

    sheets: { sheet_name: [[header_row], [data_row], ...] }
    """
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default empty sheet
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_ctx(
    blob_bytes: bytes,
    doc_id: uuid.UUID | None = None,
    row_level_rag: bool = False,
) -> tuple[BlobRef, IngestContext]:
    doc_id = doc_id or uuid.uuid4()
    blob = BlobRef(
        bucket="test",
        key="test.xlsx",
        mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        size_bytes=len(blob_bytes),
    )
    ctx = IngestContext(
        document_id=doc_id,
        tenant_id=uuid.UUID("00000000-0000-0000-0000-000000000001"),
        blob=blob,
        filename="test.xlsx",
        config={"row_level_rag": row_level_rag},
        _settings=MagicMock(),
    )
    ctx.read_blob = AsyncMock(return_value=blob_bytes)
    return blob, ctx


# ---------------------------------------------------------------------------
# Handler metadata
# ---------------------------------------------------------------------------

def test_xlsx_handler_name():
    assert XlsxHandler.name == "xlsx"


def test_xlsx_handler_accepts_ooxml_mime():
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in XlsxHandler.accepts


def test_xlsx_handler_accepts_legacy_excel_mime():
    assert "application/vnd.ms-excel" in XlsxHandler.accepts


def test_xlsx_handler_cost_class_is_io():
    assert XlsxHandler.cost_class == "io"


# ---------------------------------------------------------------------------
# Extraction — single sheet
# ---------------------------------------------------------------------------

async def test_xlsx_single_sheet_produces_one_table():
    data = _build_xlsx({"Sheet1": [["name", "age"], ["Alice", 30], ["Bob", 25]]})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert len(result.tables) == 1
    assert result.tables[0].headers == ["name", "age"]
    assert len(result.tables[0].rows) == 2


async def test_xlsx_table_rows_are_dicts():
    data = _build_xlsx({"Data": [["x", "y"], [1, 2], [3, 4]]})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.tables[0].rows[0] == {"x": "1", "y": "2"}
    assert result.tables[0].rows[1] == {"x": "3", "y": "4"}


async def test_xlsx_table_name_matches_sheet_name():
    data = _build_xlsx({"Inventory": [["item", "qty"], ["widget", 10]]})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.tables[0].name == "Inventory"
    assert result.tables[0].sheet == "Inventory"


async def test_xlsx_document_id_matches_context():
    doc_id = uuid.uuid4()
    data = _build_xlsx({"S": [["a"], [1]]})
    blob, ctx = _make_ctx(data, doc_id=doc_id)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.document_id == doc_id


async def test_xlsx_source_handler_name():
    data = _build_xlsx({"S": [["a"], [1]]})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.source_handler == "xlsx"


async def test_xlsx_metadata_populated():
    data = _build_xlsx({"Sheet1": [["col"], [1]], "Sheet2": [["col"], [2]]})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.metadata["format"] == "xlsx"
    assert result.metadata["sheet_count"] == 2
    assert "Sheet1" in result.metadata["sheets"]
    assert "Sheet2" in result.metadata["sheets"]


# ---------------------------------------------------------------------------
# Extraction — multiple sheets
# ---------------------------------------------------------------------------

async def test_xlsx_multi_sheet_produces_multiple_tables():
    data = _build_xlsx({
        "Revenue": [["month", "amount"], ["Jan", 1000], ["Feb", 1200]],
        "Costs":   [["month", "amount"], ["Jan", 800], ["Feb", 850]],
    })
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert len(result.tables) == 2
    names = {t.name for t in result.tables}
    assert names == {"Revenue", "Costs"}


async def test_xlsx_each_sheet_has_correct_rows():
    data = _build_xlsx({
        "A": [["v"], [1], [2]],
        "B": [["v"], [3]],
    })
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    by_name = {t.name: t for t in result.tables}
    assert len(by_name["A"].rows) == 2
    assert len(by_name["B"].rows) == 1


# ---------------------------------------------------------------------------
# Row-level RAG
# ---------------------------------------------------------------------------

async def test_xlsx_row_level_rag_disabled_by_default():
    """row_level_rag defaults to False — no fragments produced."""
    data = _build_xlsx({"S": [["a", "b"], [1, 2], [3, 4]]})
    blob, ctx = _make_ctx(data, row_level_rag=False)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.fragments == []
    assert result.tables  # table still produced


async def test_xlsx_row_level_rag_enabled_produces_fragments():
    data = _build_xlsx({"S": [["name", "score"], ["Alice", 95], ["Bob", 87]]})
    blob, ctx = _make_ctx(data, row_level_rag=True)
    result = await XlsxHandler().extract(blob, ctx)

    assert len(result.fragments) == 2
    assert all(f.kind == "table_row" for f in result.fragments)


async def test_xlsx_row_fragments_have_table_lineage():
    data = _build_xlsx({"S": [["k", "v"], ["a", "1"]]})
    blob, ctx = _make_ctx(data, row_level_rag=True)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.fragments[0].table_lineage is not None
    assert "table_id" in result.fragments[0].table_lineage


async def test_xlsx_row_fragment_content_uses_pipe_separator():
    data = _build_xlsx({"S": [["city", "pop"], ["Istanbul", "15M"]]})
    blob, ctx = _make_ctx(data, row_level_rag=True)
    result = await XlsxHandler().extract(blob, ctx)

    content = result.fragments[0].content
    assert "city: Istanbul" in content
    assert "pop: 15M" in content
    assert "|" in content


async def test_xlsx_row_level_rag_capped_at_500():
    """Row fragments are capped at 500 regardless of actual row count."""
    rows = [["id", "val"]] + [[i, f"v{i}"] for i in range(600)]
    data = _build_xlsx({"Big": rows})
    blob, ctx = _make_ctx(data, row_level_rag=True)
    result = await XlsxHandler().extract(blob, ctx)

    assert len(result.fragments) == 500
    assert len(result.tables[0].rows) == 600  # full table still stored


# ---------------------------------------------------------------------------
# Edge cases
# ---------------------------------------------------------------------------

async def test_xlsx_empty_sheet_produces_no_table():
    """A completely empty sheet should not produce a table."""
    data = _build_xlsx({"Empty": []})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.tables == []


async def test_xlsx_all_empty_rows_skipped():
    """Rows where every cell is empty are not included in the table."""
    data = _build_xlsx({"S": [["a", "b"], [None, None], [1, 2]]})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    # Only one non-empty data row
    assert len(result.tables[0].rows) == 1
    assert result.tables[0].rows[0]["a"] == "1"


async def test_xlsx_none_cells_become_empty_string():
    data = _build_xlsx({"S": [["a", "b"], [1, None]]})
    blob, ctx = _make_ctx(data)
    result = await XlsxHandler().extract(blob, ctx)

    assert result.tables[0].rows[0]["b"] == ""
