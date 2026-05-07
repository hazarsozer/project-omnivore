"""Unit tests for pipeline/handlers/docx.py — synthetic in-memory DOCX fixtures."""
from __future__ import annotations

import io
import uuid
from unittest.mock import AsyncMock, MagicMock

from omnivore.pipeline.context import BlobRef, IngestContext
from omnivore.pipeline.handlers.docx import DocxHandler

# ---------------------------------------------------------------------------
# Fixture helpers
# ---------------------------------------------------------------------------

def _build_docx(paragraphs: list[str], heading: str | None = None) -> bytes:
    """Build a DOCX in memory — no disk, no S3."""
    from docx import Document

    doc = Document()
    if heading:
        doc.add_heading(heading, level=1)
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_ctx(blob_bytes: bytes, doc_id: uuid.UUID | None = None) -> tuple[BlobRef, IngestContext]:
    doc_id = doc_id or uuid.uuid4()
    blob = BlobRef(
        bucket="test-bucket",
        key="test.docx",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=len(blob_bytes),
    )
    ctx = IngestContext(
        document_id=doc_id,
        tenant_id=uuid.UUID("00000000-0000-0000-0000-000000000001"),
        blob=blob,
        filename="test.docx",
        config={},
        _settings=MagicMock(),
    )
    ctx.read_blob = AsyncMock(return_value=blob_bytes)
    return blob, ctx


# ---------------------------------------------------------------------------
# Handler metadata
# ---------------------------------------------------------------------------

def test_docx_handler_name():
    assert DocxHandler.name == "docx"


def test_docx_handler_accepts_ooxml_mime():
    assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in DocxHandler.accepts


def test_docx_handler_accepts_legacy_msword_mime():
    assert "application/msword" in DocxHandler.accepts


def test_docx_handler_cost_class_is_cpu():
    assert DocxHandler.cost_class == "cpu"


# ---------------------------------------------------------------------------
# Extraction — paragraphs
# ---------------------------------------------------------------------------

async def test_docx_extracts_paragraph_content():
    content = _build_docx(["First paragraph.", "Second paragraph."])
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    texts = {f.content for f in result.fragments}
    assert "First paragraph." in texts
    assert "Second paragraph." in texts


async def test_docx_result_document_id_matches_context():
    content = _build_docx(["text"])
    doc_id = uuid.uuid4()
    blob, ctx = _make_ctx(content, doc_id=doc_id)

    result = await DocxHandler().extract(blob, ctx)

    assert result.document_id == doc_id


async def test_docx_source_handler_is_set():
    content = _build_docx(["text"])
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    assert result.source_handler == "docx"


async def test_docx_metadata_format_is_docx():
    content = _build_docx(["text"])
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    assert result.metadata.get("format") == "docx"


# ---------------------------------------------------------------------------
# Extraction — headings + heading_path propagation
# ---------------------------------------------------------------------------

async def test_docx_heading_present_in_blocks():
    content = _build_docx(["Body text."], heading="Introduction")
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    kinds = {b.kind for b in result.blocks}
    assert "heading" in kinds


async def test_docx_paragraph_heading_path_includes_heading():
    content = _build_docx(["Body text."], heading="Introduction")
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    para_fragments = [f for f in result.fragments if f.kind == "text"]
    assert len(para_fragments) >= 1
    assert "Introduction" in para_fragments[0].heading_path


async def test_docx_no_heading_means_empty_heading_path():
    content = _build_docx(["Just a paragraph."])
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    assert result.fragments[0].heading_path == []


# ---------------------------------------------------------------------------
# Extraction — empty document
# ---------------------------------------------------------------------------

async def test_docx_empty_document_no_fragments():
    content = _build_docx([])
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    assert result.fragments == []


async def test_docx_empty_document_no_blocks():
    content = _build_docx([])
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    assert result.blocks == []


# ---------------------------------------------------------------------------
# Extraction — tables
# ---------------------------------------------------------------------------

async def test_docx_table_extracted_with_correct_headers():
    from docx import Document as DocxDocument

    doc = DocxDocument()
    table = doc.add_table(rows=3, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "foo"
    table.rows[1].cells[1].text = "bar"
    table.rows[2].cells[0].text = "baz"
    table.rows[2].cells[1].text = "qux"
    buf = io.BytesIO()
    doc.save(buf)
    blob, ctx = _make_ctx(buf.getvalue())

    result = await DocxHandler().extract(blob, ctx)

    assert len(result.tables) == 1
    assert result.tables[0].headers == ["Name", "Value"]


async def test_docx_table_rows_are_dicts():
    from docx import Document as DocxDocument

    doc = DocxDocument()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Col1"
    table.rows[0].cells[1].text = "Col2"
    table.rows[1].cells[0].text = "a"
    table.rows[1].cells[1].text = "b"
    buf = io.BytesIO()
    doc.save(buf)
    blob, ctx = _make_ctx(buf.getvalue())

    result = await DocxHandler().extract(blob, ctx)

    assert result.tables[0].rows == [{"Col1": "a", "Col2": "b"}]


async def test_docx_table_name_is_prefixed():
    from docx import Document as DocxDocument

    doc = DocxDocument()
    t = doc.add_table(rows=2, cols=1)
    t.rows[0].cells[0].text = "Header"
    t.rows[1].cells[0].text = "Val"
    buf = io.BytesIO()
    doc.save(buf)
    blob, ctx = _make_ctx(buf.getvalue())

    result = await DocxHandler().extract(blob, ctx)

    assert result.tables[0].name.startswith("table_")


# ---------------------------------------------------------------------------
# Extraction — blocks are created for every non-empty paragraph
# ---------------------------------------------------------------------------

async def test_docx_blocks_created_for_paragraphs_and_headings():
    content = _build_docx(["Para one.", "Para two."], heading="Chapter")
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    kinds = [b.kind for b in result.blocks]
    assert "heading" in kinds
    assert "paragraph" in kinds


async def test_docx_fragment_source_block_ids_non_empty():
    content = _build_docx(["Some content."])
    blob, ctx = _make_ctx(content)

    result = await DocxHandler().extract(blob, ctx)

    for frag in result.fragments:
        assert len(frag.source_block_ids) > 0


async def test_docx_whitespace_only_paragraph_is_skipped():
    """Paragraphs whose text.strip() is empty must not produce a block or fragment."""
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_paragraph("   ")  # whitespace only — should be skipped
    doc.add_paragraph("Real content.")
    buf = io.BytesIO()
    doc.save(buf)
    blob, ctx = _make_ctx(buf.getvalue())

    result = await DocxHandler().extract(blob, ctx)

    assert len(result.fragments) == 1
    assert result.fragments[0].content == "Real content."


